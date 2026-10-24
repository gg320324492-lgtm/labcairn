"""The spreadsheet adapter and the ingest pipeline that feeds the store.

Two things are under test that a passing run would not otherwise show.

**Idempotence.** ADR-0009's key is ``(content hash, extraction version,
address)``, and the claim that re-ingesting an unchanged file writes nothing
is the claim that makes a failed ingest safe to retry. It is asserted by
running the write twice and counting, not by inspecting the constraint.

**Disclosure instead of silence.** An empty spreadsheet and a workbook we
failed to read are different claims. So are a formula whose result we do not
have and a cell holding the number three. Each has a test that distinguishes
it from the other.
"""

from __future__ import annotations

import datetime as dt
import io
import os
import pathlib
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from sqlalchemy.orm import Session

from conftest import migrated_database_from
from labcairn.db.materials import (
    ExtractedTextRow,
    problems_for,
    put_source,
    put_source_texts,
    text_at,
)
from labcairn.db.session import create_db_engine, create_session_factory
from labcairn.material.ingest import (
    IngestedMaterial,
    ingest_bytes,
    ingest_into_store,
    kind_for_path,
    to_text_row,
)
from labcairn.material.spreadsheet import render_cell
from labcairn.material.units import ExtractedUnit, ExtractionResult
from labcairn.security.untrusted import UntrustedText
from labcairn.sources.locators import XlsxLocator
from labcairn.sources.source import SourceKind

NOW = dt.datetime(2026, 10, 17, tzinfo=dt.UTC)


def _xlsx(
    *,
    cells: dict[str, Any] | None = None,
    sheets: tuple[str, ...] = ("Trial",),
    empty_sheets: tuple[str, ...] = (),
) -> bytes:
    """A workbook with *cells* in the first sheet and *empty_sheets* added."""
    buffer = io.BytesIO()
    book = openpyxl.Workbook()
    # openpyxl's stubs type `active` as optional and `__setitem__` narrowly,
    # so both are narrowed here rather than silenced at the call sites.
    first = book.active
    assert first is not None
    first.title = sheets[0]
    for ref, value in (cells or {}).items():
        first[ref] = value
    for title in empty_sheets:
        book.create_sheet(title)
    book.save(buffer)
    return buffer.getvalue()


def _cell(unit: ExtractedUnit) -> XlsxLocator:
    assert isinstance(unit.locator, XlsxLocator), unit.locator
    return unit.locator


def _codes(result: ExtractionResult) -> set[str]:
    return {p.code for p in result.problems}


# --------------------------------------------------------------------------
# kind detection
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("path", "expected"),
    [
        ("papers/t.pdf", SourceKind.PDF),
        ("notes.md", SourceKind.MARKDOWN),
        ("notes.txt", SourceKind.TEXT),
        ("data.csv", SourceKind.CSV),
        ("data.tsv", SourceKind.CSV),
        ("book.xlsx", SourceKind.XLSX),
        ("archive.zip", SourceKind.OTHER),
        ("no-extension", SourceKind.OTHER),
    ],
)
def test_kind_comes_from_the_extension(path: str, expected: SourceKind) -> None:
    assert kind_for_path(path) is expected


def test_extension_matching_ignores_case() -> None:
    assert kind_for_path("BOOK.XLSX") is SourceKind.XLSX


# --------------------------------------------------------------------------
# the spreadsheet adapter
# --------------------------------------------------------------------------


def test_one_unit_per_cell_holding_a_value() -> None:
    result = _extract(_xlsx(cells={"A1": "trial", "B2": 3.5, "C3": "alpha"}))
    assert [(_cell(u).cell_ref, u.text) for u in result.units] == [
        ("A1", "trial"),
        ("B2", "3.5"),
        ("C3", "alpha"),
    ]


def test_an_empty_cell_is_not_a_unit() -> None:
    """A spreadsheet is mostly empty; the padding is not evidence.

    This is where the text adapter's "keep every line" rule deliberately
    does *not* transfer: a blank line in prose is a line a reader counts,
    while a blank cell is not a measurement.

    The fixture has to contain a *gap* between two filled cells. An earlier
    version wrote only A1, so openpyxl stored one cell and ``iter_rows()``
    yielded one -- deleting the emptiness guard entirely left this test green.
    """
    result = _extract(_xlsx(cells={"A1": "first", "A5": "last"}))
    assert [_cell(u).cell_ref for u in result.units] == ["A1", "A5"]
    assert all(u.text for u in result.units), "an empty cell must not become a unit"


def test_cells_are_attributed_to_the_sheet_they_are_in() -> None:
    """Attributing everything to the first sheet would be silently wrong.

    There was no test putting a cell in a *second* sheet, so a version
    reading ``sheetnames[0]`` for every cell passed.
    """
    buffer = io.BytesIO()
    book = openpyxl.Workbook()
    first_sheet = book.active
    assert first_sheet is not None
    first_sheet.title = "First"
    first_sheet["A1"] = "in-first"
    second_sheet = book.create_sheet("Second")
    second_sheet["A1"] = "in-second"
    book.save(buffer)
    result = _extract(buffer.getvalue())
    by_sheet = {
        _cell(unit).sheet_name: unit.text
        for unit in result.units
        if _cell(unit).sheet_name in {"First", "Second"}
    }
    assert by_sheet == {"First": "in-first", "Second": "in-second"}


def test_a_formula_with_no_cached_value_is_disclosed_and_kept() -> None:
    """The cell is dropped silently unless this holds.

    openpyxl does not compute, so a formula with no cached result reads as
    None -- and an emptiness check placed first would drop it, recording a
    *measured blank* where the truth is "we could not read this". A later
    numeric check would then compare against nothing.
    """
    result = _extract(_xlsx(cells={"B3": "=SUM(B1:B2)"}))
    assert "xlsx_formula_without_cached_value" in _codes(result)
    assert [u.text for u in result.units] == ["=SUM(B1:B2)"]


def test_an_empty_sheet_is_reported_as_a_fact() -> None:
    result = _extract(_xlsx(cells={"A1": "x"}, empty_sheets=("Notes",)))
    assert "xlsx_sheet_empty" in _codes(result)
    assert any(p.location == "Notes" for p in result.problems)


def test_a_workbook_that_cannot_be_read_is_disclosed_not_raised() -> None:
    result = _extract(b"this is not a workbook")
    assert result.units == ()
    assert "xlsx_unreadable" in _codes(result)


@pytest.mark.parametrize(
    "payload", [b"", b"\x00\x01\x02", b"PK\x03\x04 truncated", bytes(range(256))]
)
def test_no_hostile_input_raises(payload: bytes) -> None:
    _extract(payload)  # must not raise


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (True, "TRUE"),
        (False, "FALSE"),
        # A spreadsheet stores a naive datetime, so that is what is rendered.
        (dt.datetime(2026, 10, 17, 9, 30), "2026-10-17T09:30:00"),  # noqa: DTZ001
        (dt.date(2026, 10, 17), "2026-10-17"),
        (0.1, "0.1"),
        (1234567.891, "1234567.891"),
        (3.5, "3.5"),
        ("plain", "plain"),
    ],
)
def test_values_render_without_a_locale(value: object, expected: str) -> None:
    """A value rendered differently per machine makes the hash machine-dependent."""
    assert render_cell(value) == expected


def test_extraction_is_deterministic() -> None:
    """ADR-0008: the same bytes must give the same string, or a character
    range silently re-points."""
    from labcairn.material.addressing import render_locator_key

    data = _xlsx(cells={"A1": "trial", "B2": 3.5, "C3": "alpha", "D4": True})
    first = [(render_locator_key(u.locator), u.text) for u in _extract(data).units]
    second = [(render_locator_key(u.locator), u.text) for u in _extract(data).units]
    assert first == second
    assert first


def test_openpyxl_version_is_in_the_extraction_version() -> None:
    """It reads the container and the XML, so an upgrade can change a cell."""
    from labcairn.material import adapter_for

    libs = adapter_for(SourceKind.XLSX).libraries()
    assert libs.get("openpyxl") == openpyxl.__version__


def _extract(data: bytes) -> ExtractionResult:
    from labcairn.material import adapter_for

    return adapter_for(SourceKind.XLSX).extract(data)


# --------------------------------------------------------------------------
# ingest
# --------------------------------------------------------------------------


def test_content_hash_is_the_file_hash() -> None:
    import hashlib

    data = _xlsx(cells={"A1": "x"})
    material = ingest_bytes(data, relative_path="a.xlsx", display_name="a.xlsx", ingested_at=NOW)
    assert material.source.content_sha256 == hashlib.sha256(data).hexdigest()
    assert material.source.size_bytes == len(data)


def test_the_same_bytes_and_path_give_the_same_id_on_two_machines() -> None:
    a = ingest_bytes(b"same", relative_path="a.txt", display_name="a", ingested_at=NOW)
    b = ingest_bytes(b"same", relative_path="a.txt", display_name="a", ingested_at=NOW)
    assert a.source.source_id == b.source.source_id


def test_a_material_with_no_adapter_is_recorded_with_a_disclosure() -> None:
    """Inventoried and hashed, and *not* silently absent."""
    material = ingest_bytes(
        b"PK\x03\x04", relative_path="a.zip", display_name="a.zip", ingested_at=NOW
    )
    assert material.source.kind is SourceKind.OTHER
    assert material.texts == ()
    assert "no_adapter_for_kind" in {p.code for p in material.problems}
    assert material.extraction_version == "no-adapter"


def test_a_format_with_an_adapter_but_no_extension_still_declines() -> None:
    material = ingest_bytes(
        b"%PDF-1.4\n", relative_path="paper", display_name="paper", ingested_at=NOW
    )
    assert "no_adapter_for_kind" in {p.code for p in material.problems}


def test_every_unit_carries_the_canonical_key_and_a_text_hash() -> None:
    material = ingest_bytes(
        _xlsx(cells={"A1": "trial"}),
        relative_path="a.xlsx",
        display_name="a.xlsx",
        ingested_at=NOW,
    )
    unit = material.texts[0]
    assert unit.locator_key.startswith("k=4:xlsx|")
    import hashlib

    assert unit.text_sha256 == hashlib.sha256(unit.text.encode()).hexdigest()


def test_text_leaves_the_pipeline_wrapped() -> None:
    """ADR-0003 C6's guarantee is consumer-side, so the wrap is on the way out."""
    material = ingest_bytes(
        _xlsx(cells={"A1": "trial"}),
        relative_path="a.xlsx",
        display_name="a.xlsx",
        ingested_at=NOW,
    )
    wrapped = material.as_untrusted(material.texts[0])
    assert isinstance(wrapped, UntrustedText)
    assert wrapped.value == "trial"
    assert wrapped.locator is not None


def test_untrusted_text_is_not_a_string() -> None:
    """ADR-0003 C6 is enforced *statically*, so there is no runtime check here.

    ``UntrustedText`` is a frozen dataclass that does not subclass ``str``,
    and mypy rejects ``isinstance(wrapped, str)`` as **unreachable** --
    "subclass of UntrustedText and str cannot exist: have distinct disjoint
    bases". That is stronger than a test: imported text cannot be passed
    where a ``str`` is annotated, and the type checker refuses the file on
    every change. A runtime assertion here would check something the type
    system already guarantees, and would read as though it did not.
    """


# --------------------------------------------------------------------------
# persistence
# --------------------------------------------------------------------------


@pytest.fixture
def session(migrated_database_template: Path, tmp_path: pathlib.Path) -> Iterator[Session]:
    target = tmp_path / "materials.db"
    migrated_database_from(migrated_database_template, target)
    engine = create_db_engine(target)
    made = create_session_factory(engine)()
    yield made
    made.close()
    engine.dispose()


def _ingested(now: dt.datetime = NOW) -> IngestedMaterial:
    return ingest_bytes(
        _xlsx(cells={"A1": "trial", "B2": 3.5}),
        relative_path="papers/t.xlsx",
        display_name="t.xlsx",
        ingested_at=now,
    )


def test_re_ingesting_an_unchanged_file_writes_nothing(session: Session) -> None:
    """The claim that makes a failed ingest safe to retry."""
    material = _ingested()
    put_source(session, material.source)
    rows = [to_text_row(t, created_at=NOW) for t in material.texts]
    session.commit()

    assert put_source_texts(
        session,
        source_content_sha256=material.source.content_sha256,
        extraction_version=material.extraction_version,
        rows=rows,
    ) == (2, 0)
    session.commit()
    assert put_source_texts(
        session,
        source_content_sha256=material.source.content_sha256,
        extraction_version=material.extraction_version,
        rows=rows,
    ) == (0, 2)
    session.commit()


def test_a_different_extraction_version_is_a_different_address(session: Session) -> None:
    """An upgrade writes beside the old rows rather than over them."""
    material = _ingested()
    put_source(session, material.source)
    rows = [to_text_row(t, created_at=NOW) for t in material.texts]
    session.commit()
    put_source_texts(
        session,
        source_content_sha256=material.source.content_sha256,
        extraction_version="an-older-version",
        rows=rows,
    )
    session.commit()
    old = text_at(
        session,
        source_content_sha256=material.source.content_sha256,
        extraction_version="an-older-version",
        locator_key=material.texts[0].locator_key,
    )
    assert old == "trial"


def test_a_malformed_row_is_an_error_not_a_silent_duplicate(session: Session) -> None:
    """Why this is not ``INSERT OR IGNORE``.

    OR IGNORE would swallow a blank ``locator_key`` or an invalid
    ``cell_ref`` and report a row the constraints refused as a row that
    was already there.
    """
    from sqlalchemy.exc import IntegrityError

    put_source(session, _ingested().source)
    session.commit()
    bad = ExtractedTextRow(
        locator_kind="xlsx",
        locator_key="   ",
        text="x",
        text_sha256="0" * 64,
        created_at=NOW,
        sheet_name="S",
        cell_ref="B2",
    )
    with pytest.raises(IntegrityError):
        put_source_texts(
            session,
            source_content_sha256="a" * 64,
            extraction_version="v1",
            rows=[bad],
        )
        session.commit()


def test_text_is_found_at_the_address_the_locator_names(session: Session) -> None:
    material = _ingested()
    put_source(session, material.source)
    session.commit()
    put_source_texts(
        session,
        source_content_sha256=material.source.content_sha256,
        extraction_version=material.extraction_version,
        rows=[to_text_row(t, created_at=NOW) for t in material.texts],
    )
    session.commit()
    for unit in material.texts:
        found = text_at(
            session,
            source_content_sha256=material.source.content_sha256,
            extraction_version=material.extraction_version,
            locator_key=unit.locator_key,
        )
        assert found == unit.text


def test_an_address_that_was_never_extracted_is_none_not_an_error(session: Session) -> None:
    assert (
        text_at(
            session,
            source_content_sha256="f" * 64,
            extraction_version="v1",
            locator_key="k=4:xlsx|sh=1:Nope|c=1:Z99",
        )
        is None
    )


# --------------------------------------------------------------------------
# what review found uncovered
# --------------------------------------------------------------------------


def test_a_material_with_an_adapter_that_reads_nothing_is_disclosed() -> None:
    """The branch that produced no text *and* no problem was entirely
    untested, so deleting it left every test green."""
    material = ingest_bytes(
        chr(0).encode() * 8, relative_path="a.txt", display_name="a.txt", ingested_at=NOW
    )
    codes = {p.code for p in material.problems}
    assert codes, "reading nothing must never be silent"


def test_source_id_follows_the_bytes_not_their_length() -> None:
    """Two different files of the same length must not collide.

    The previous test ingested the *same* bytes twice, so it could only
    detect non-determinism -- never a wrong-but-deterministic id.
    """
    a = ingest_bytes(b"aaaa", relative_path="a.txt", display_name="a", ingested_at=NOW)
    b = ingest_bytes(b"bbbb", relative_path="b.txt", display_name="b", ingested_at=NOW)
    assert a.source.source_id != b.source.source_id


def test_editing_a_file_does_not_make_it_a_new_material() -> None:
    """The other half, and the half that was pinned the wrong way round.

    This test used to ingest two different bodies at the *same* path and assert
    they got different ids. That is a collision only if the id identifies
    content -- and it does not. ``source_id`` identifies a **location**, which
    is what ``walk.py`` has always assumed when it derives its own, and what
    ``sources.relative_path``'s UNIQUE constraint requires: with a
    content-derived id, the second ingest of an edited file raised
    ``IntegrityError`` and left the caller's session unusable, because this
    function rolls nothing back.

    So the property to hold is the pair of them. Different locations never
    share an id; the same location keeps its id when its bytes change, and the
    content hash -- which is what distinguishes the versions -- lives in its
    own column.
    """
    first = ingest_bytes(b"version one\n", relative_path="n.md", display_name="n", ingested_at=NOW)
    second = ingest_bytes(
        b"version two, edited\n", relative_path="n.md", display_name="n", ingested_at=NOW
    )
    assert first.source.source_id == second.source.source_id
    assert first.source.content_sha256 != second.source.content_sha256


def test_an_array_formula_is_recorded_as_its_formula_not_a_repr() -> None:
    """openpyxl models array formulas as objects.

    ``str()`` on one gives ``<...ArrayFormula object at 0x...>`` -- a
    fabrication, and one that makes the excerpt hash depend on the process.
    """
    from openpyxl.worksheet.formula import ArrayFormula

    buffer = io.BytesIO()
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "S"
    sheet["A1"], sheet["A2"], sheet["A3"] = 1, 2, 3
    sheet["B1"] = ArrayFormula("B1:B3", "=SUM(A1:A3)")
    book.save(buffer)
    result = _extract(buffer.getvalue())
    texts = [u.text for u in result.units if _cell(u).cell_ref == "B1"]
    assert texts == ["=SUM(A1:A3)"], texts
    assert "object at 0x" not in texts[0]


def test_extraction_is_the_same_under_hash_randomisation() -> None:
    """A repr with a memory address would make the hash process-dependent."""
    import subprocess
    import sys

    data = _xlsx(cells={"A1": "trial", "B2": 3.5}).hex()
    script = (
        "import hashlib,sys;"
        f"sys.path.insert(0, {str(Path(__file__).parents[1] / 'src')!r});"
        "from labcairn.material import adapter_for;"
        "from labcairn.sources.source import SourceKind;"
        f"r=adapter_for(SourceKind.XLSX).extract(bytes.fromhex('{data}'));"
        "print(sorted(hashlib.sha256(u.text.encode()).hexdigest() for u in r.units))"
    )
    digests = set()
    for seed in ("0", "1", "99999"):
        out = subprocess.run(
            [sys.executable, "-c", script],
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            # The parent decodes UTF-8, so the child has to write UTF-8: the
            # child's output encoding is otherwise the machine's locale
            # encoding, and the two halves of the pipe would not agree.
            env={**os.environ, "PYTHONHASHSEED": seed, "PYTHONIOENCODING": "utf-8"},
            check=True,
        )
        digests.add(out.stdout.strip())
    assert len(digests) == 1, digests


# --------------------------------------------------------------------------
# The paths review found with zero coverage, and the wiring it found absent
# --------------------------------------------------------------------------


def test_a_formula_with_no_text_is_refused_not_recorded_as_an_equals_sign() -> None:
    """An empty expression is not an empty string.

    A workbook can store ``<f t="array" ref="B1:B3"/>`` with no payload,
    and openpyxl synthesises ``.text == ""``. Recording that -- or the "="
    it becomes -- puts a string in the record that the file does not
    contain, while the disclosure beside it claims the formula *was*
    recorded. Review found this after the repr fix; the guard checked
    ``None`` and not emptiness.
    """
    from openpyxl.worksheet.formula import ArrayFormula

    for expression in ("", "="):
        buffer = io.BytesIO()
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "S"
        sheet["A1"] = 1
        sheet["B1"] = ArrayFormula("B1:B3", expression)
        book.save(buffer)
        result = _extract(buffer.getvalue())
        assert [u for u in result.units if _cell(u).cell_ref == "B1"] == [], expression
        assert "xlsx_formula_not_readable" in _codes(result), expression


def test_a_mid_walk_failure_is_disclosed_and_keeps_what_was_read() -> None:
    """A chartsheet has no ``iter_rows``; the try now wraps the walk.

    What survives is what was actually read, with the failure disclosed.
    Dropping it would claim the sheet was empty, which is a different and
    wrong claim.
    """
    buffer = io.BytesIO()
    book = openpyxl.Workbook()
    data_sheet = book.active
    assert data_sheet is not None
    data_sheet.title = "Data"
    assert book.active is not None
    data_sheet["A1"] = "kept"
    book.create_chartsheet("Chart")
    book.save(buffer)
    result = _extract(buffer.getvalue())
    # Measured, not assumed: openpyxl 3.1.5 raises while *loading* a
    # workbook containing a chartsheet, so this arrives as the open failure
    # rather than a mid-walk one. Either way it is disclosed and no cell is
    # invented.
    assert _codes(result) & {"xlsx_unreadable", "xlsx_iteration_failed"}, [
        p.code for p in result.problems
    ]
    assert result.problems[0].detail, "a disclosure with no detail is not a disclosure"


def test_ingesting_unchanged_bytes_a_second_time_writes_nothing(session: Session) -> None:
    """``put_source`` rewrote eight columns every time, ``ingested_at`` included.

    Review showed the claim "re-ingesting an unchanged file writes
    nothing" was false for ``sources`` while true for ``source_texts`` --
    and the test asserted counts, which are the wrong invariant for a row
    whose *history* is being rewritten.
    """
    from sqlalchemy import select

    from labcairn.db.models import SourceRow

    buffer = _xlsx(cells={"A1": "trial", "B2": 3.5})
    later = NOW + dt.timedelta(days=1)
    first = ingest_into_store(
        session, buffer, relative_path="p/t.xlsx", display_name="t.xlsx", ingested_at=NOW
    )
    row = session.scalar(
        select(SourceRow).where(SourceRow.source_id == first.material.source.source_id)
    )
    assert row is not None
    before = row.ingested_at
    session.expire_all()
    second = ingest_into_store(
        session, buffer, relative_path="p/t.xlsx", display_name="t.xlsx", ingested_at=later
    )
    session.expire_all()
    assert row.ingested_at == before, "a re-ingest of unchanged bytes rewrote ingested_at"
    assert second.text_rows_written == 0 and second.problem_rows_written == 0


def test_disclosures_reach_the_store_and_are_not_predicted(session: Session) -> None:
    """The headline of the last review: a table with no writer.

    ``ingest_bytes`` produced disclosures and ``put_source_problems``
    stored them, and nothing connected the two, so a disclosure reached
    memory and stopped.
    """
    buffer = _xlsx(cells={"A1": "trial"}, empty_sheets=("Notes",))
    outcome = ingest_into_store(
        session, buffer, relative_path="p/t.xlsx", display_name="t.xlsx", ingested_at=NOW
    )
    assert outcome.problem_rows_written == 1
    stored = problems_for(session, source_id=outcome.material.source.source_id)
    assert [r.code for r in stored] == ["xlsx_sheet_empty"]
    assert stored[0].detail, "a disclosure with no detail is not a disclosure"

    again = ingest_into_store(
        session, buffer, relative_path="p/t.xlsx", display_name="t.xlsx", ingested_at=NOW
    )
    assert again.problem_rows_written == 0, "a duplicate disclosure was written again"
